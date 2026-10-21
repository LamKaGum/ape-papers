#!/usr/bin/env python3
"""把省赛原始名单整理成 整理/{年份}/{省份}/{赛道}.csv。"""

from __future__ import annotations

import csv
import hashlib
import re
import subprocess
import zipfile
from collections import defaultdict
from pathlib import Path

ROOT = Path("/workspace/china-iec-awards")
OUT = ROOT / "整理"
CACHE = Path("/tmp/iec-convert")
FIELDS = [
    "年份",
    "省份",
    "赛道",
    "奖项",
    "组别",
    "项目名称",
    "学校",
    "负责人",
    "成员",
    "指导教师",
    "来源文件",
    "备注",
]
TRACK_ORDER = ["高教主赛道", "红旅赛道", "职教赛道", "产业赛道", "萌芽赛道", "其他", "未分赛道"]
AWARD_ORDER = [
    "冠军",
    "亚军",
    "季军",
    "金奖",
    "银奖",
    "铜奖",
    "一等奖",
    "二等奖",
    "三等奖",
    "优秀奖",
    "创新潜力奖",
    "优胜奖",
]
HEADER_KEYS = [
    ("项目名称", ["项目名称", "参赛项目", "命题名称"]),
    ("学校", ["学校名称", "所在学校", "推荐学校", "申报高校", "所属学校", "学校"]),
    ("负责人", ["项目负责人", "负责人姓名", "团队负责人", "负责人", "姓名"]),
    ("成员", ["团队成员姓名", "团队负责人及成员", "团队成员", "项目成员", "参赛成员", "参赛队员", "成员"]),
    ("指导教师", ["指导教师", "指导老师", "指教教师", "优秀创新创业导师"]),
    ("奖项", ["拟授奖项", "所获奖项", "授予奖项", "获奖等次", "拟获奖励", "拟授奖", "奖项"]),
    ("赛道", ["参赛赛道", "赛道类型", "赛道"]),
    ("组别", ["参赛组别", "赛道名称", "组别", "组名"]),
    ("附注", ["单项奖"]),
]


def norm_text(value) -> str:
    if value is None:
        return ""
    text = str(value).replace("\u3000", "").replace("\xa0", "")
    text = text.replace("\r", "").replace("\n", "")
    text = re.sub(r"\s+", "", text)
    return text.strip()


def key_norm(value) -> str:
    text = norm_text(value)
    return re.sub(r"[“”\"'‘’—–\-·.。,，、:：;；（）()【】\[\]《》]", "", text)


def clean_cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).replace("\u3000", " ").replace("\xa0", " ")
    text = text.replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n+", "", text).strip()
    return text


def tracks_in(text: str) -> list[str]:
    raw = norm_text(text)
    found = []
    if "萌芽" in raw:
        found.append("萌芽赛道")
    if "职教" in raw:
        found.append("职教赛道")
    if "产业" in raw:
        found.append("产业赛道")
    if "红旅" in raw or ("红" in raw and "筑梦" in raw) or ("青年红色" in raw):
        found.append("红旅赛道")
    if "高教" in raw or "主赛道" in raw:
        found.append("高教主赛道")
    # 主赛道常与高教同时出现，保留高教主赛道一次
    dedup = []
    for item in found:
        if item not in dedup:
            dedup.append(item)
    return dedup


def sole_track(text: str) -> str:
    found = tracks_in(text)
    if len(found) == 1:
        return found[0]
    return ""


def normalize_award(text: str) -> str:
    raw = norm_text(text)
    raw = re.sub(r"^[一二三四五六七八九十0-9]+[、.．]", "", raw)
    raw = raw.replace("拟授", "").replace("拟获", "")
    match = re.match(
        r"^(冠军|亚军|季军|金奖|银奖|铜奖|一等奖|二等奖|三等奖|优秀奖|创新潜力奖|优胜奖)",
        raw,
    )
    if match and len(raw) <= 24:
        return match.group(1)
    if raw.endswith("奖") and 2 <= len(raw) <= 14 and "名单" not in raw and "赛道" not in raw:
        return raw
    return ""


def file_note(name: str) -> str:
    notes = []
    if "拟" in name:
        notes.append("拟获奖或拟授奖")
    if "复赛" in name:
        notes.append("复赛成绩，不是正式公布稿")
    if "入围" in name:
        notes.append("含现场赛入围，不是完整金银奖名单")
    if "更正" in name:
        notes.append("项目信息更正")
    if "决赛" in name and "获奖" not in name and "成绩" not in name:
        notes.append("决赛名单，原文未必标注奖项等次")
    if "成绩" in name:
        notes.append("成绩公示")
    return "；".join(notes)


def locked_track(name: str) -> str:
    if any(word in name for word in ("单项", "导师", "指导教师", "组织奖", "组织校", "优秀组织", "工作者", "优胜")):
        return "其他"
    return ""


def default_award(name: str) -> str:
    if "导师" in name or "指导教师" in name:
        return "优秀指导教师"
    if "组织" in name:
        return "优秀组织奖"
    if "工作者" in name:
        return "年度创新创业教育工作者"
    if "优胜" in name:
        return "优胜奖"
    return ""


def map_header(cells: list[str]) -> dict[str, int] | None:
    best: dict[str, tuple[int, int]] = {}
    for index, cell in enumerate(cells):
        folded = norm_text(cell)
        if not folded or len(folded) > 20:
            continue
        for field, keys in HEADER_KEYS:
            for key in keys:
                if folded == key and (field not in best or len(key) > best[field][0]):
                    best[field] = (len(key), index)
    chosen = {field: pos for field, (_length, pos) in best.items()}
    aliases = {
        "参赛组别/赛道": "赛道",
        "组别/赛道": "赛道",
        "赛道/组别": "赛道",
        "所属学校": "学校",
        "所在学校": "学校",
    }
    for index, cell in enumerate(cells):
        folded = norm_text(cell)
        field = aliases.get(folded)
        if field and field not in best:
            best[field] = (8, index)
    chosen = {field: pos for field, (_length, pos) in best.items()}
    strong = [field for field in ("项目名称", "学校", "负责人", "成员", "指导教师", "奖项", "赛道") if field in chosen]
    if len(strong) >= 2 and ("项目名称" in chosen or "学校" in chosen):
        return chosen
    if "学校" in chosen and any(norm_text(cell) == "序号" for cell in cells):
        return chosen
    return None


def is_header_repeat(cells: list[str], colmap: dict[str, int]) -> bool:
    if "项目名称" not in colmap:
        return False
    value = norm_text(cells[colmap["项目名称"]]) if colmap["项目名称"] < len(cells) else ""
    return value in {"项目名称", "参赛项目", "命题名称"}


def org_header(cells: list[str]) -> list[str] | None:
    """只有整行都是赛道名称时，才当成优秀组织奖的分赛道表头。"""
    nonempty = [cell for cell in cells if norm_text(cell)]
    if len(nonempty) < 2:
        return None
    labels = []
    for cell in cells:
        folded = norm_text(cell)
        if not folded:
            labels.append("")
            continue
        track = sole_track(folded)
        pure = bool(track) and "赛道" in folded and len(folded) <= 20
        pure = pure and not any(token in folded for token in ("获奖", "名单", "项目", "奖", "组"))
        labels.append(track if pure else "")
    named = [item for item in labels if item]
    if len(named) >= 2 and len(set(named)) >= 2 and len(named) == len(nonempty):
        return labels
    return None


class Parser:
    def __init__(self, year: str, province: str, source: str, note: str):
        self.year = year
        self.province = province
        self.source = source
        self.note = note
        self.file_locked = locked_track(source)
        self.track = self.file_locked or sole_track(source)
        self.track_locked = bool(self.file_locked)
        self.award = default_award(source)
        self.colmap: dict[str, int] | None = None
        self.org_cols: list[str] | None = None
        self.section_note = ""
        self.rows: list[dict] = []

    def set_track(self, text: str) -> None:
        if self.track_locked:
            return
        found = tracks_in(text)
        if len(found) == 1:
            self.track = found[0]

    def set_award(self, text: str) -> None:
        award = normalize_award(text)
        if award:
            self.award = award

    def consider_banner(self, text: str) -> bool:
        folded = norm_text(text)
        if not folded or len(folded) > 80:
            return False
        award = normalize_award(folded)
        found = tracks_in(folded)
        if award and (len(folded) <= 16 or folded.startswith(award)):
            self.award = award
            return True
        if re.search(r"\d项", folded) and len(folded) > 12:
            return False
        if len(found) >= 2 and len(folded) <= 40:
            if not self.track_locked:
                self.track = ""
            self.section_note = "原文合并公布：" + "、".join(found)
            return True
        track_tail = folded.endswith(("职教", "高教", "产业", "萌芽", "筑梦", "主赛道"))
        if len(found) == 1 and len(folded) <= 70 and ("赛道" in folded or (track_tail and len(folded) <= 36)):
            self.section_note = ""
            self.set_track(folded)
            if self.track_locked:
                self.section_note = found[0]
            if not normalize_award(folded):
                self.award = default_award(self.source)
            return True
        if folded in {"金奖", "银奖", "铜奖", "一等奖", "二等奖", "三等奖", "优秀奖", "冠军", "亚军", "季军"}:
            self.award = folded
            return True
        return False

    def emit(self, data: dict) -> None:
        project = data.get("项目名称", "")
        school = data.get("学校", "")
        leader = data.get("负责人", "")
        members = data.get("成员", "")
        if norm_text(project) in {
            "序号", "项目名称", "参赛项目", "命题名称", "无", "负责人", "团队成员",
            "指导老师", "指导教师", "赛道", "组别", "奖项", "学校", "学校名称",
        }:
            return
        if len(project) > 180 or len(school) > 40:
            return
        if "名单" in project and len(project) > 20:
            return
        if not project and not (school and (leader or members or data.get("奖项"))):
            return
        if school and norm_text(school).isdigit():
            return
        row = {field: "" for field in FIELDS}
        row.update(data)
        row["年份"] = self.year
        row["省份"] = self.province
        row["来源文件"] = self.source
        row["赛道"] = row["赛道"] or "未分赛道"
        extra = self.note
        if row["备注"] and extra:
            row["备注"] = row["备注"] + "；" + extra
        elif extra:
            row["备注"] = extra
        self.rows.append(row)

    def consume(self, matrix: list[list[str]]) -> None:
        for raw in matrix:
            cells = [clean_cell(cell) for cell in raw]
            while cells and not cells[-1]:
                cells.pop()
            if not any(cells):
                continue
            unique = []
            for cell in cells:
                if cell and cell not in unique:
                    unique.append(cell)
            if len(unique) == 1 and self.consider_banner(unique[0]):
                self.org_cols = None
                continue
            header = map_header(cells)
            if header:
                self.colmap = header
                self.org_cols = None
                continue
            labels = org_header(cells)
            if labels:
                self.org_cols = labels
                self.colmap = None
                continue
            if self.org_cols:
                for cell, track in zip(cells, self.org_cols):
                    if not cell or not track or cell == track:
                        continue
                    self.emit(
                        {
                            "赛道": "其他",
                            "奖项": self.award or "优秀组织奖",
                            "学校": cell,
                            "备注": f"原赛道：{track}",
                        }
                    )
                continue
            if not self.colmap or is_header_repeat(cells, self.colmap):
                continue
            data = {
                "赛道": self.track,
                "奖项": self.award,
                "组别": "",
                "项目名称": "",
                "学校": "",
                "负责人": "",
                "成员": "",
                "指导教师": "",
                "备注": "",
            }
            for field, index in self.colmap.items():
                if index >= len(cells):
                    continue
                value = cells[index]
                if field == "赛道":
                    picked = sole_track(value)
                    if picked:
                        data["赛道"] = picked
                        if not self.track_locked:
                            self.track = picked
                    parts = [part for part in value.replace("//", "/").split("/") if part]
                    group = next((part for part in parts if part and not sole_track(part)), "")
                    if group and not data["组别"]:
                        data["组别"] = group
                    continue
                if field == "奖项":
                    award = normalize_award(value) or value
                    if award and award not in {"奖项", "拟授奖项"}:
                        data["奖项"] = award
                    continue
                if field == "附注":
                    if value:
                        data["备注"] = f"单项奖：{value}"
                    continue
                if field == "组别":
                    picked = sole_track(value)
                    if picked and not data["赛道"]:
                        data["赛道"] = picked
                    group = value
                    if "/" in value:
                        left, _right = value.split("/", 1)
                        if sole_track(value):
                            data["赛道"] = sole_track(value)
                            group = left
                    data["组别"] = group
                    continue
                data[field] = value
            if self.track_locked or any(token in data["奖项"] for token in ("组织", "导师", "集体", "工作者")):
                data["赛道"] = "其他"
            if not data["赛道"] and self.section_note:
                data["备注"] = (data["备注"] + "；" if data["备注"] else "") + self.section_note
            self.emit(data)


def read_xlsx(path: Path) -> list[tuple[str, list[list[str]]]]:
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    sheets = []
    try:
        for name in book.sheetnames:
            sheet = book[name]
            rows = []
            for row in sheet.iter_rows(values_only=True):
                rows.append(["" if cell is None else cell for cell in row])
            if sheet_is_stats(name, rows):
                continue
            sheets.append((name, rows))
    finally:
        book.close()
    return sheets


def read_xls(path: Path) -> list[tuple[str, list[list[str]]]]:
    import xlrd

    book = xlrd.open_workbook(path)
    sheets = []
    for name in book.sheet_names():
        sheet = book.sheet_by_name(name)
        rows = []
        for r in range(sheet.nrows):
            rows.append([sheet.cell_value(r, c) for c in range(sheet.ncols)])
        if sheet_is_stats(name, rows):
            continue
        sheets.append((name, rows))
    return sheets


def sheet_is_stats(name: str, rows: list[list]) -> bool:
    """跳过只有金银铜计数的汇总表，保留表名带“汇总”但实际是名单的工作表。"""
    if "汇总" not in name or "名单" in name:
        return False
    for row in rows[:12]:
        cells = [clean_cell(cell) for cell in row]
        if map_header(cells):
            return False
    return True


def parse_workbook(path: Path, year: str, province: str) -> list[dict]:
    suffix = path.suffix.lower()
    sheets = read_xlsx(path) if suffix in {".xlsx", ".xlsm"} else read_xls(path)
    try:
        source = str(path.relative_to(ROOT))
    except ValueError:
        source = path.name
    parser = Parser(year, province, source, file_note(path.name))
    for name, rows in sheets:
        parser.track_locked = bool(parser.file_locked or locked_track(name))
        if parser.track_locked:
            parser.track = "其他"
        else:
            parser.track = sole_track(name) or sole_track(path.name)
        parser.award = normalize_award(name) or default_award(name) or default_award(path.name)
        parser.section_note = ""
        parser.colmap = None
        parser.org_cols = None
        parser.consume(rows)
    return parser.rows


def docx_blocks(path: Path):
    from docx import Document
    from docx.oxml.ns import qn

    document = Document(path)
    body = document.element.body
    table_iter = iter(document.tables)
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            text = "".join(node.text or "" for node in child.iter(qn("w:t")))
            yield ("p", text)
        elif child.tag == qn("w:tbl"):
            table = next(table_iter)
            matrix = []
            for row in table.rows:
                matrix.append([cell.text for cell in row.cells])
            yield ("t", matrix)


def parse_docx(path: Path, year: str, province: str, source_name: str) -> list[dict]:
    parser = Parser(year, province, source_name, file_note(Path(source_name).name))
    blocks = list(docx_blocks(path))
    has_table = any(kind == "t" for kind, _payload in blocks)
    for kind, payload in blocks:
        if kind == "p":
            if parser.consider_banner(payload):
                continue
            if has_table:
                continue
            text = clean_cell(payload)
            folded = norm_text(text)
            if not (2 <= len(folded) <= 30):
                continue
            if any(
                token in folded
                for token in ("附件", "名单", "大赛", "通知", "公布", "关于", "互联网", "建行", "大学生", "创新创业", "赛道")
            ):
                continue
            note = f"原赛道：{parser.section_note}" if parser.section_note else ""
            parser.emit(
                {
                    "赛道": parser.track or "其他",
                    "奖项": parser.award,
                    "学校": text,
                    "备注": note,
                }
            )
        else:
            parser.consume(payload)
    return parser.rows


def convert_office(src: Path) -> Path:
    CACHE.mkdir(parents=True, exist_ok=True)
    digest = hashlib.md5(src.read_bytes()).hexdigest()[:12]
    target = CACHE / f"{digest}.docx"
    if target.exists() and target.stat().st_size > 0:
        return target
    subprocess.run(
        [
            "soffice",
            "--headless",
            "--norestore",
            "--convert-to",
            "docx",
            "--outdir",
            str(CACHE),
            str(src),
        ],
        check=True,
        timeout=180,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    produced = CACHE / (src.stem + ".docx")
    if produced.exists() and produced != target:
        produced.replace(target)
    if not target.exists():
        raise FileNotFoundError(f"未转出 docx：{src}")
    return target


def matrix_from_table(table) -> list[list[str]]:
    """超过 15 列通常是跨行表格被拆碎，不写入分赛道表。"""
    extracted = table.extract() or []
    if extracted and len(extracted[0]) > 15:
        return []
    return extracted


def parse_pdf(path: Path, year: str, province: str) -> list[dict]:
    import pdfplumber

    try:
        source = str(path.relative_to(ROOT))
    except ValueError:
        source = path.name
    parser = Parser(year, province, source, file_note(path.name))
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            tables = page.find_tables() or []
            words = page.extract_words() or []
            events = []
            for table in tables:
                x0, y0, x1, y1 = table.bbox
                events.append((y0, 1, "t", matrix_from_table(table)))
            lines: dict[int, list] = defaultdict(list)
            for word in words:
                cx = (word["x0"] + word["x1"]) / 2
                cy = (word["top"] + word["bottom"]) / 2
                inside = False
                for table in tables:
                    b = table.bbox
                    if b[0] - 1 <= cx <= b[2] + 1 and b[1] - 1 <= cy <= b[3] + 1:
                        inside = True
                        break
                if inside:
                    continue
                lines[round(word["top"])].append(word)
            for top, group in lines.items():
                group.sort(key=lambda item: item["x0"])
                text = "".join(item["text"] for item in group)
                events.append((top, 0, "p", text))
            for _top, _order, kind, payload in sorted(events, key=lambda item: (item[0], item[1])):
                if kind == "p":
                    parser.consider_banner(payload)
                else:
                    parser.consume(payload)
    return parser.rows


def rel_source(path: Path) -> str:
    try:
        return str(path.relative_to(ROOT))
    except ValueError:
        return path.name


def parse_any(path: Path, year: str, province: str, source_name: str | None = None) -> list[dict]:
    source = source_name or rel_source(path)
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm", ".xls"}:
        rows = parse_workbook(path, year, province)
        for row in rows:
            row["来源文件"] = source
        return rows
    if suffix == ".docx":
        return parse_docx(path, year, province, source)
    if suffix in {".doc", ".wps"}:
        converted = convert_office(path)
        return parse_docx(converted, year, province, source)
    if suffix == ".pdf":
        rows = parse_pdf(path, year, province)
        for row in rows:
            row["来源文件"] = source
        return rows
    return []


def extract_guangdong() -> list[tuple[Path, str]]:
    archive = ROOT / "2023" / "广东" / "广东省分赛获奖名单及广州萌芽名单.zip"
    if not archive.exists():
        return []
    dest = CACHE / "guangdong"
    dest.mkdir(parents=True, exist_ok=True)
    found = []
    with zipfile.ZipFile(archive) as zipped:
        for info in zipped.infolist():
            try:
                name = info.filename.encode("cp437").decode("gbk")
            except UnicodeError:
                name = info.filename
            target = dest / Path(name).name
            target.write_bytes(zipped.read(info))
            found.append((target, f"2023/广东/{name}"))
    return found


def iter_sources():
    seen_hash = {}
    files = []
    for path in sorted(ROOT.rglob("*")):
        if not path.is_file():
            continue
        if any(part in { "整理", "tools" } for part in path.parts):
            continue
        if path.suffix.lower() not in {".pdf", ".doc", ".docx", ".wps", ".xls", ".xlsx", ".xlsm"}:
            continue
        if "2025" in path.parts and "广西" in path.parts and path.suffix.lower() == ".xlsx":
            continue
        if path.name.endswith(".pdf.xlsx"):
            # 江西这份与同目录 PDF 重复，表格本身仍可用；PDF 另计
            pass
        digest = hashlib.md5(path.read_bytes()).hexdigest()
        if digest in seen_hash:
            continue
        seen_hash[digest] = path
        year = path.relative_to(ROOT).parts[0]
        province = path.relative_to(ROOT).parts[1]
        if year not in {"2023", "2024", "2025"}:
            continue
        files.append((path, year, province, None))
    for path, source in extract_guangdong():
        files.append((path, "2023", "广东", source))
    return files


def priority(row: dict) -> tuple:
    name = row["来源文件"]
    score = 0
    if name.endswith((".xlsx", ".xls", ".xlsm")):
        score += 30
    elif name.endswith((".doc", ".docx", ".wps")):
        score += 22
    elif name.endswith(".pdf"):
        score += 12
    if "正式" in name:
        score += 24
    if "拟" in name:
        score -= 4
    filled = sum(1 for field in ("奖项", "组别", "项目名称", "学校", "负责人", "成员", "指导教师") if row.get(field))
    return (score, filled, len(row.get("成员", "")), len(row.get("项目名称", "")))


def dedup(rows: list[dict]) -> list[dict]:
    grouped: dict[tuple, list[dict]] = defaultdict(list)
    for row in rows:
        ident = (
            row["年份"],
            row["省份"],
            row["赛道"],
            key_norm(row["项目名称"])[:80],
            key_norm(row["学校"]),
            key_norm(row["负责人"]),
        )
        grouped[ident].append(row)
    kept: list[dict] = []
    for items in grouped.values():
        by_award: dict[str, dict] = {}
        for item in items:
            award = key_norm(item["奖项"])
            current = by_award.get(award)
            if current is None or priority(item) > priority(current):
                by_award[award] = item
        if "" in by_award and len(by_award) > 1:
            del by_award[""]
        kept.extend(by_award.values())
    return kept


def award_rank(name: str) -> int:
    try:
        return AWARD_ORDER.index(name)
    except ValueError:
        return len(AWARD_ORDER) + 1


def write_outputs(rows: list[dict], log_lines: list[str]) -> None:
    if OUT.exists():
        for old in OUT.rglob("*"):
            if old.is_file():
                old.unlink()
    grouped: dict[tuple, list] = defaultdict(list)
    for row in rows:
        grouped[(row["年份"], row["省份"], row["赛道"])].append(row)
    for (year, province, track), items in grouped.items():
        folder = OUT / year / province
        folder.mkdir(parents=True, exist_ok=True)
        target = folder / f"{track}.csv"
        with target.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, FIELDS)
            writer.writeheader()
            items.sort(key=lambda item: (award_rank(item["奖项"]), item["项目名称"]))
            writer.writerows(items)
    stats = OUT / "统计.csv"
    with stats.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["年份", "省份", "赛道", "条数"])
        for key in sorted(grouped, key=lambda item: (item[0], item[1], TRACK_ORDER.index(item[2]) if item[2] in TRACK_ORDER else 99)):
            writer.writerow([key[0], key[1], key[2], len(grouped[key])])
    log = OUT / "_解析日志.tsv"
    log.write_text("\n".join(log_lines) + "\n", encoding="utf-8")


def parse_job(job: tuple) -> tuple:
    path, year, province, source = job
    label = source or rel_source(path)
    try:
        rows = parse_any(path, year, province, source)
        status = "ok"
    except Exception as exc:  # noqa: BLE001
        rows = []
        status = f"失败:{type(exc).__name__}:{exc}"[:180]
    return label, year, province, rows, status


def main() -> None:
    from concurrent.futures import ProcessPoolExecutor, as_completed

    log = ["年份\t省份\t文件\t行数\t状态"]
    collected = []
    jobs = iter_sources()
    office = [job for job in jobs if Path(job[0]).suffix.lower() != ".pdf"]
    pdfs = [job for job in jobs if Path(job[0]).suffix.lower() == ".pdf"]
    print(f"表格与文档 {len(office)}，PDF {len(pdfs)}", flush=True)
    for job in office:
        label, year, province, rows, status = parse_job(job)
        print(f"{len(rows):5d}  {label}  {status if status != 'ok' else ''}", flush=True)
        log.append(f"{year}\t{province}\t{label}\t{len(rows)}\t{status}")
        collected.extend(rows)
    if pdfs:
        with ProcessPoolExecutor(max_workers=4) as pool:
            futures = [pool.submit(parse_job, job) for job in pdfs]
            for future in as_completed(futures):
                label, year, province, rows, status = future.result()
                print(f"{len(rows):5d}  {label}  {status if status != 'ok' else ''}", flush=True)
                log.append(f"{year}\t{province}\t{label}\t{len(rows)}\t{status}")
                collected.extend(rows)
    merged = dedup(collected)
    print(f"去重前 {len(collected)} 去重后 {len(merged)}", flush=True)
    write_outputs(merged, log)


if __name__ == "__main__":
    main()
